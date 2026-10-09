from openpyxl import Workbook
from openpyxl.worksheet.table import Table
from openpyxl.worksheet.table import TableStyleInfo


def exporter_excel(df, fichier_sortie):

    wb = Workbook()

    ws = wb.active
    ws.title = "Factures"

    # Entêtes
    for col, nom_colonne in enumerate(df.columns, start=1):
        ws.cell(
            row=1,
            column=col,
            value=nom_colonne
        )

    # Données
    for ligne, row in enumerate(
        df.itertuples(index=False),
        start=2
    ):
        for col, valeur in enumerate(row, start=1):
            ws.cell(
                row=ligne,
                column=col,
                value=valeur
            )

    # Création du tableau Excel
    derniere_ligne = ws.max_row
    derniere_colonne = ws.max_column

    lettre_fin = ws.cell(
        1,
        derniere_colonne
    ).column_letter

    table = Table(
        displayName="Tableau1",
        ref=f"A1:{lettre_fin}{derniere_ligne}"
    )

    style = TableStyleInfo(
        name="TableStyleLight8",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False
    )

    table.tableStyleInfo = style

    ws.add_table(table)

    wb.save(fichier_sortie)